"""Reproducible, offline-only SaaS planning workbook and independent numeric snapshot."""

import argparse
import copy
import datetime as dt
import json
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUTPUT = ROOT / "documents" / "saas-plan"
SCENARIOS = ("Lean", "Base", "Growth")
CATEGORIES = ("INFRA", "TOOLS", "LEGAL", "SECURITY", "COMPLIANCE", "INSURANCE",
              "ACCOUNTING", "SUPPORT", "MARKETING", "ONBOARD", "AI", "CTMS")
MONEY = '"$"#,##0;[Red]("$"#,##0);"$"0'
NUMBER = '#,##0.00;[Red](#,##0.00)'
BLUE = "DCEEFF"
GREEN = "E5F2E2"
AMBER = "FFF0C2"
NAVY = "17365D"


def load_inputs():
    return json.loads((HERE / "inputs.json").read_text())


def month_date(start, month):
    date = dt.date.fromisoformat(start) if isinstance(start, str) else start
    index = date.year * 12 + date.month - 1 + month - 1
    return dt.date(index // 12, index % 12 + 1, 1)


def assumptions(data, scenario):
    col = SCENARIOS.index(scenario) + 2
    return {row[0]: row[col] for row in data["assumptions"]}


def staff_fte(staff, scenario, month, a, controls):
    start, end, *ramp = staff["ramps"][scenario]
    if staff["scope"] != "Core":
        if not controls[staff["scope"].lower()]:
            return 0
        start = int(a[staff["scope"].lower() + "_start"])
    if not start <= month <= end:
        return 0
    return ramp[0 if month <= 6 else 1 if month <= 15 else 2 if month <= 30 else 3]


def evaluate(data, scenario):
    """Numeric companion; formulas in Excel use the same input rows and conventions."""
    a = assumptions(data, scenario)
    ctrl = data["controls"]
    results = []
    cash = a["initial_cash"]
    cash_no_revenue = cash
    customers = cumulative = need = 0.0
    for month in range(1, 61):
        year = (month - 1) // 12
        new = (a["new_customers"] * (1 + a["customer_growth"]) **
               ((month - a["launch"]) // 12)) if month >= a["launch"] else 0
        customers = customers * (1 - a["churn"]) + new
        row = {
            "month": month, "date": str(month_date(data["start"], month)),
            "new": new, "customers": customers, "users": customers * a["users"],
            "requests_m": customers * a["users"] * a["requests"] * 30 / 1e6,
            "storage": customers * a["storage"], "logs": customers * a["logs"],
            "fte": 0, "heads": 0, "labor": 0, "recruiting": 0, "equipment": 0,
            "unpaid": 0, "quote": 0,
            **{cat: 0 for cat in CATEGORIES}
        }
        for staff in data["staff"]:
            fte = staff_fte(staff, scenario, month, a, ctrl)
            previous = staff_fte(staff, scenario, month - 1, a, ctrl) if month > 1 else 0
            salary = a["founder_salary"] if staff["type"] == "Founder" else staff["salary"]
            economic = fte * salary / 12 * (1 + a["salary_growth"]) ** year
            paid = staff["type"] != "Founder" or ctrl["founder_paid"]
            row["fte"] += fte
            row["heads"] += math.ceil(fte)
            row["labor"] += economic * (1 + (a["burden"] if staff["type"] == "Employee" else 0)) if paid else 0
            row["unpaid"] += economic if not paid else 0
            row["recruiting"] += max(0, fte - previous) * a["recruiting"] if staff["type"] == "Employee" else 0
            row["equipment"] += max(0, fte - previous) * a["equipment"]
        quantities = {
            "Fixed": 1, "Seats": math.ceil(row["fte"]), "Customers": customers,
            "NewCustomers": new, "StorageGB": row["storage"], "LogGB": row["logs"],
            "RequestsM": row["requests_m"], "InputM": customers * a["tokens_in"] / 1e6,
            "OutputM": customers * a["tokens_out"] / 1e6
        }
        for cost in data["costs"]:
            _, category, _, stack, scope, start, end, frequency, driver, rate, _, status, _ = cost
            if stack not in ("Both", ctrl["stack"]) or (scope != "Core" and not ctrl[scope.lower()]):
                continue
            start = max(start, int(a[scope.lower() + "_start"])) if scope != "Core" else start
            if not start <= month <= end:
                continue
            occurs = frequency == "Monthly" or (frequency == "Once" and month == start) or (
                frequency == "Annual" and (month - start) % 12 == 0)
            amount = quantities[driver] * rate * (1 + a["inflation"]) ** year if occurs else 0
            row[category] += amount
            if status in ("Quote required", "Unverified"):
                row["quote"] += amount
        row["vendor"] = sum(row[cat] for cat in CATEGORIES)
        row["contingency"] = (row["labor"] + row["recruiting"] + row["equipment"] + row["vendor"]) * a["contingency"]
        row["revenue"] = customers * a["arpa"]
        row["payment"] = row["revenue"] * a["payment_rate"]
        row["spend"] = row["labor"] + row["recruiting"] + row["equipment"] + row["vendor"] + row["contingency"] + row["payment"]
        row["arr"] = row["revenue"] * 12
        row["burn"] = row["spend"] - row["revenue"]
        row["financing"] = a["financing_amount"] if month == a["financing_month"] else 0
        row["opening"] = cash
        cash += row["financing"] - row["burn"]
        cash_no_revenue += row["financing"] - row["spend"]
        cumulative += row["spend"]
        need += row["burn"]
        row.update(closing=cash, cumulative=cumulative, gap=max(0, -cash),
                   no_revenue=cash_no_revenue, need=need,
                   optional=row["AI"] + row["CTMS"])
        results.append(row)
    annual = []
    for year in range(1, 6):
        rows = results[(year - 1) * 12: year * 12]
        annual.append({"year": year, "spend": sum(r["spend"] for r in rows),
                       "revenue": sum(r["revenue"] for r in rows),
                       "closing": rows[-1]["closing"]})
    calendar_years = []
    for year in sorted({r["date"][:4] for r in results}):
        rows = [r for r in results if r["date"][:4] == year]
        calendar_years.append({"year": int(year), "spend": sum(r["spend"] for r in rows),
                               "revenue": sum(r["revenue"] for r in rows)})
    return {
        "scenario": scenario, "spend": sum(r["spend"] for r in results),
        "year1": annual[0]["spend"], "peak_fte": max(r["fte"] for r in results),
        "cash_low": min(r["closing"] for r in results),
        "gap": max(r["gap"] for r in results),
        "no_revenue_gap": max(0, -min(r["no_revenue"] for r in results)),
        "quote_subtotal": sum(r["quote"] for r in results),
        "unpaid_founder": sum(r["unpaid"] for r in results),
        "optional_vendor": sum(r["optional"] for r in results),
        "first_negative": ("Not funded" if a["initial_cash"] <= 0 and results[0]["financing"] <= 0 and results[0]["revenue"] <= 0
                           else next((r["month"] for r in results if r["closing"] < 0), ">60 months")),
        "annual": annual, "calendar": calendar_years, "monthly": results
    }


class Planner:
    def __init__(self, data, scope, sources):
        self.data, self.scope, self.sources = data, scope, sources
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        self.wb.properties.creator = "Clinical Roundup planning generator"
        self.wb.properties.title = "Clinical Roundup — SaaS Build and Scale"
        self.wb.properties.description = "Synthetic planning assumptions; model, not quote."
        self.ar = {r[0]: i + 6 for i, r in enumerate(data["assumptions"])}
        self.staff_rows = {}
        self.cost_rows = {}

    def sheet(self, name, title, note, headers=None):
        ws = self.wb.create_sheet(name)
        ws.append([title])
        ws.append([note])
        ws.cell(3, 1, "← Start Here").hyperlink = "#'Start Here'!A1"
        ws.cell(3, 1).style = "Hyperlink"
        ws.row_dimensions[1].height = 29
        ws.row_dimensions[2].height = 44
        ws.cell(1, 1).font = Font(size=17, bold=True, color=NAVY)
        ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
        if headers:
            for col, value in enumerate(headers, 1):
                ws.cell(5, col, value)
                ws.cell(5, col).fill = PatternFill("solid", fgColor=NAVY)
                ws.cell(5, col).font = Font(bold=True, color="FFFFFF")
                ws.cell(5, col).alignment = Alignment(wrap_text=True)
            ws.row_dimensions[5].height = 32
        ws.freeze_panes = "C6"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws

    @staticmethod
    def dropdown(ws, cells, options):
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"')
        dv.errorTitle = "Choose an allowed value"
        dv.error = "Use the dropdown list."
        dv.showErrorMessage = True
        dv.errorStyle = "stop"
        ws.add_data_validation(dv)
        dv.add(cells)

    @staticmethod
    def bound(ws, cells, low, high, kind="decimal"):
        dv = DataValidation(type=kind, operator="between", formula1=str(low), formula2=str(high))
        dv.showErrorMessage = True
        dv.error = f"Enter a value between {low} and {high}."
        ws.add_data_validation(dv)
        dv.add(cells)

    def a(self, key, scenario):
        col = "CDE"[SCENARIOS.index(scenario)]
        return f"'Assumptions'!${col}${self.ar[key]}"

    @staticmethod
    def scope_formula(scope):
        return {"Core": "1", "AI": "'Start Here'!$B$7", "CTMS": "'Start Here'!$B$8"}[scope]

    def setup(self):
        ws = self.sheet("Start Here", "Clinical Roundup | Option C build and scale",
                        "MODEL, NOT QUOTE. USD cash planning; US hiring benchmark; 60 months. Blue = editable, green = formula, amber = evidence/quote gap. No PHI or live PoC changes.")
        controls = [
            ("Scenario", self.data["selected"]), ("Project start (first day of month)", dt.date.fromisoformat(self.data["start"])),
            ("Enable optional AI (0/1)", self.data["controls"]["ai"]),
            ("Enable optional CTMS (0/1)", self.data["controls"]["ctms"]),
            ("Mutually exclusive cloud candidate", self.data["controls"]["stack"]),
            ("Pay founder cash (0/1)", self.data["controls"]["founder_paid"])
        ]
        for r, (label, value) in enumerate(controls, 5):
            ws.cell(r, 1, label)
            ws.cell(r, 2, value).fill = PatternFill("solid", fgColor=BLUE)
        ws["B6"].number_format = "yyyy-mm-dd"
        self.dropdown(ws, "B5", SCENARIOS)
        self.dropdown(ws, "B9", ("Azure", "AWS"))
        self.bound(ws, "B7:B8", 0, 1, "whole")
        self.bound(ws, "B10", 0, 1, "whole")
        date_validation = DataValidation(type="custom", formula1="DAY(B6)=1")
        date_validation.showErrorMessage = True
        ws.add_data_validation(date_validation)
        date_validation.add("B6")
        ws["A12"] = "Decision metrics (after Excel recalculation; see snapshot.json for offline results)"
        metrics = [
            ("5-year total cash spend", "=SUM('Monthly Model'!AC6:AC65)"),
            ("Year 1 spend / cost-only burn", "=SUM('Monthly Model'!AC6:AC17)"),
            ("Peak FTE incl. founder", "=MAX('Monthly Model'!I6:I65)"),
            ("Cash low point with revenue/financing", "=MIN('Monthly Model'!AI6:AI65)"),
            ("Maximum financing gap (not sum of deficits)", "=MAX('Monthly Model'!AK6:AK65)"),
            ("Cost-only financing gap (zero revenue)", "=MAX(0,-MIN('Monthly Model'!AL6:AL65))"),
            ("Quote/unverified vendor allowance subtotal", "=SUM('Monthly Model'!AM6:AM65)"),
            ("Unresolved selected catalog line count", f"=SUM('Nonlabor'!N6:N{5+len(self.data['costs'])})"),
            ("Unpaid founder opportunity cost (outside spend)", "=SUM('Monthly Model'!AO6:AO65)"),
            ("Runway: first negative month / >60 / not funded",
             '=IF(AND(INDEX(\'Assumptions\'!C6:E6,1,MATCH(B5,\'Assumptions\'!C5:E5,0))<=0,\'Monthly Model\'!AG6<=0,\'Monthly Model\'!AD6<=0),"Not funded",IF(COUNTIF(\'Monthly Model\'!AI6:AI65,"<0")=0,">60 months",MATCH(1,\'Monthly Model\'!AR6:AR65,0)))')
        ]
        for r, (label, formula) in enumerate(metrics, 13):
            ws.cell(r, 1, label)
            ws.cell(r, 2, formula).number_format = NUMBER if r in (15, 20, 22) else MONEY
        ws["A24"] = "Use: 1 choose scenario/start/options; 2 review blue inputs; 3 resolve amber sources; 4 inspect monthly/annual and gates."
        ws["A25"] = "No guarantee of PMF, valuation, revenue, compliance or enterprise readiness. Negative cash is not funded."
        ws["A26"] = "Baseline = clinical rounding. CTMS requires new study/site workflows, integrations and validation; AI has separate approval gates."
        ws["A27"] = "Ranges: low/medium confidence estimates. Cash excludes income/sales taxes, debt interest, equity, depreciation and working-capital delays."
        ws["A28"] = "Scenario changes do not prove faster delivery. Gates override calendar; stop before unapproved PHI, AI or research processing."
        ws.column_dimensions["A"].width = 63
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["D"].width = 37
        ws.column_dimensions["E"].width = 28

        ws = self.sheet("Assumptions", "Editable scenario assumptions", "Selected values are formulas; scenario columns are independent assumptions, not total-cost multipliers. Start and switches live on Start Here.",
                        ["Key", "Assumption", *SCENARIOS, "Unit", "Basis / caveat", "Selected"])
        for row in self.data["assumptions"]:
            key, label, lean, base, growth, unit, note = row
            r = ws.max_row + 1
            ws.append([key, label, lean, base, growth, unit, note,
                       f'=INDEX(C{r}:E{r},1,MATCH(\'Start Here\'!$B$5,$C$5:$E$5,0))'])
            for col in range(3, 6):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=BLUE)
                ws.cell(r, col).number_format = "0.0%" if unit == "percent" else MONEY if unit.startswith("USD") else NUMBER
            ws.cell(r, 8).number_format = ws.cell(r, 3).number_format
            self.bound(ws, f"C{r}:E{r}", 0, 1 if unit == "percent" else 60 if unit == "month" else 1e9)
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["G"].width = 74

    def catalogs(self):
        ws = self.sheet("Sources", "Cost catalog and public source register",
                        "Research access/attempt date is recorded, not inferred. Unavailable primary pages remain unverified. Selected allowances are not vendor quotes; public wage surveys are not startup hiring rates.",
                        ["Source ID", "Publisher", "Product / role", "Exact public URL", "Access / attempt", "Effective / survey", "Currency", "Region", "Unit", "Public benchmark", "Selected allowance", "Evidence status", "Confidence", "Rationale", "Exclusions", "Refresh by"])
        fields = ["id", "publisher", "product", "url", "accessed", "effective", "currency",
                  "region", "unit", "benchmark", "selected", "status", "confidence", "rationale", "exclusions", "refresh"]
        for source in self.sources:
            ws.append([source.get(f) for f in fields])
            r = ws.max_row
            if source.get("url"):
                ws.cell(r, 4).hyperlink = source["url"]
                ws.cell(r, 4).style = "Hyperlink"
            if source["status"] != "Published price":
                ws.cell(r, 12).fill = PatternFill("solid", fgColor=AMBER)
        ws.column_dimensions["D"].width = 65
        ws.column_dimensions["N"].width = 80
        ws.column_dimensions["O"].width = 65

        ws = self.sheet("Staffing", "Phased staffing | annual selected salaries and fractional capacity",
                        "FTE bands are months 1–6 / 7–15 / 16–30 / 31–60. Headcount slots ≠ FTE or hiring promise. Contractor annual rates are selected retainer equivalents; no employee burden. Founder cash defaults off.",
                        ["Scenario", "Role ID", "Role / responsibility", "Engagement", "Scope", "Start month", "End month", "FTE M1–6", "FTE M7–15", "FTE M16–30", "FTE M31–60", "Selected annual USD", "Benchmark ID", "Employer burden / wages", "Recruit / added FTE", "Equipment / added FTE"])
        for scenario in SCENARIOS:
            for staff in self.data["staff"]:
                r = ws.max_row + 1
                start, end, *ramp = staff["ramps"][scenario]
                ws.append([scenario, staff["id"], staff["role"], staff["type"], staff["scope"], start, end, *ramp,
                           staff["salary"], staff["source"],
                           f'={self.a("burden", scenario)}' if staff["type"] == "Employee" else 0,
                           f'={self.a("recruiting", scenario)}' if staff["type"] == "Employee" else 0,
                           f'={self.a("equipment", scenario)}'])
                self.staff_rows[(scenario, staff["id"])] = r
                for col in range(6, 13):
                    ws.cell(r, col).fill = PatternFill("solid", fgColor=BLUE)
                    ws.cell(r, col).number_format = MONEY if col == 12 else NUMBER
                ws.cell(r, 14).number_format = "0%"
                ws.cell(r, 15).number_format = ws.cell(r, 16).number_format = MONEY
        self.bound(ws, f"F6:G{ws.max_row}", 1, 60, "whole")
        self.bound(ws, f"H6:K{ws.max_row}", 0, 100)
        self.bound(ws, f"L6:L{ws.max_row}", 0, 1e7)
        ws.column_dimensions["C"].width = 54
        ws.column_dimensions["D"].width = 26

        ws = self.sheet("Nonlabor", "Non-labor cash catalog | stable cost IDs",
                        "Only one cloud stack charged. Annual items recur start+12n; Once items occur once. Optional items start at MAX(catalog start, scenario expansion start). Salary is only in Staffing; WBS labor is informational.",
                        ["Cost ID", "Category", "Line item", "Stack", "Scope", "Start", "End", "Frequency", "Quantity driver", "Selected rate USD", "Source ID", "Evidence status", "Caveat / exclusions", "Active unresolved line"])
        for cost in self.data["costs"]:
            r = ws.max_row + 1
            ws.append(cost)
            self.cost_rows[cost[0]] = r
            ws.cell(r, 14, f'=IF(AND(OR(D{r}="Both",D{r}=\'Start Here\'!$B$9),IF(E{r}="Core",1,IF(E{r}="AI",\'Start Here\'!$B$7,\'Start Here\'!$B$8))=1,OR(L{r}="Quote required",L{r}="Unverified")),1,0)')
            for col in (6, 7, 8, 9, 10):
                ws.cell(r, col).fill = PatternFill("solid", fgColor=BLUE)
            ws.cell(r, 10).number_format = MONEY
            if cost[11] in ("Quote required", "Unverified"):
                ws.cell(r, 12).fill = PatternFill("solid", fgColor=AMBER)
        self.dropdown(ws, f"H6:H{ws.max_row}", ("Monthly", "Annual", "Once"))
        self.dropdown(ws, f"I6:I{ws.max_row}", ("Fixed", "Seats", "Customers", "NewCustomers", "StorageGB", "LogGB", "RequestsM", "InputM", "OutputM"))
        self.bound(ws, f"J6:J{ws.max_row}", 0, 1e9)
        self.bound(ws, f"F6:G{ws.max_row}", 1, 60, "whole")
        ws.column_dimensions["C"].width = 58
        ws.column_dimensions["M"].width = 86

    def traceability(self):
        specs = [
            ("Requirements", "Requirements and evidence register", "Repository presence = Implemented, not Verified or In customer use. Successful M365 PoC is user-reported; no deployment verification performed.",
             "requirements", ["id", "heading", "item", "evidence", "status", "caveat", "requirement", "acceptance", "priority", "phase", "owner", "dependency", "scope", "wbs", "cost"]),
            ("Migration WBS", "Migration and refactoring WBS", "Effort is an assumed person-day range, not a separate engineering charge. STAFF references the staffing pool; no WBS labor added to financial totals. Freeze/write-pause cutover; no unqualified dual-write.",
             "wbs", ["id", "requirement", "work", "decision", "evidence", "effort_days", "start", "end", "owner", "dependency", "scope", "cost", "acceptance"]),
            ("Architecture", "Architecture candidates | managed modular monolith", "Azure East US is a modeled candidate, not mandated. AWS US East managed alternative is mutually exclusive. All HA/retention/SKUs require sizing, eligibility review and quotes; no Kubernetes prerequisite.",
             "architecture", ["id", "component", "azure", "alternative", "decision", "driver", "sku", "region", "allowance", "ha", "retention", "trigger", "effort_days", "cost"]),
            ("Roadmap", "Methodology, roadmap and decision gates", "Assume 2-week iterations, small vertical slices, acceptance criteria, PR review + tests, threat model, staged release, feedback. Gates can delay dates; added headcount does not halve timelines.",
             "roadmap", ["id", "phase", "start", "end", "dependency", "owner", "acceptance", "confidence", "scope", "risk"]),
            ("Risks and Funding", "Risks, decisions and funding gates", "Qualified specialists must assess HIPAA/BAA eligibility, clinical safety, intended use and research validation/e-signature rules. SOX is not assumed. No cloud product makes the application compliant.",
             "risks", ["id", "risk", "owner", "mitigation", "gate", "cost", "likelihood"])
        ]
        for name, title, note, key, fields in specs:
            ws = self.sheet(name, title, note, [f.replace("_", " ").title() for f in fields])
            for entry in self.scope[key]:
                ws.append([entry.get(f, "") for f in fields])
                ws.row_dimensions[ws.max_row].height = 100
            for col in range(1, len(fields) + 1):
                ws.column_dimensions[letter(col)].width = 23 if col == 1 else 52
            if name == "Requirements":
                self.dropdown(ws, f"E6:E{ws.max_row}", ("Idea", "Implemented", "Verified", "In customer use"))
            if name == "Migration WBS":
                ws.cell(5, len(fields) + 1, "Low assumed days (−30%)")
                ws.cell(5, len(fields) + 2, "High assumed days (+50%)")
                for r in range(6, ws.max_row + 1):
                    ws.cell(r, len(fields) + 1, f"=F{r}*0.7")
                    ws.cell(r, len(fields) + 2, f"=F{r}*1.5")
                    ws.cell(r, 6).fill = PatternFill("solid", fgColor=BLUE)
            if name == "Roadmap":
                for m in range(1, 61):
                    col = len(fields) + m
                    ws.cell(5, col, m)
                    ws.column_dimensions[letter(col)].width = 4
                    for r in range(6, ws.max_row + 1):
                        start = f"C{r}"
                        optional_start = f'IF(I{r}="AI",\'Assumptions\'!$H${self.ar["ai_start"]},\'Assumptions\'!$H${self.ar["ctms_start"]})'
                        effective_start = f'IF(I{r}="Core",{start},{optional_start})'
                        effective_end = f'IF(I{r}="Core",D{r},{optional_start}+D{r}-C{r})'
                        enabled = f'IF(I{r}="Core",1,IF(I{r}="AI",\'Start Here\'!$B$7,\'Start Here\'!$B$8))'
                        ws.cell(r, col, f'=IF(AND({letter(col)}$5>={effective_start},{letter(col)}$5<={effective_end},{enabled}=1),1,0)').number_format = '0;-0;;'
                    ws.conditional_formatting.add(f"{letter(col)}6:{letter(col)}{ws.max_row}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="80B7DB")))
        self.funding()

    def funding(self):
        ws = self.wb["Risks and Funding"]
        r = ws.max_row + 3
        ws.cell(r, 1, "Funding by core stage (selected scenario; spend includes contingency)")
        ws.append(["Stage", "From", "Through", "Stage gross spend", "Cumulative gross spend", "Cumulative net need before cash/financing", "Maximum gap to date"])
        for label, start, end in [("Assessment", 1, 3), ("Foundation", 4, 8), ("Synthetic pilot", 9, 11),
                                  ("Approved pilot", 12, 15), ("Repeatable delivery", 16, 24), ("Growth", 25, 60)]:
            ws.append([label, start, end,
                       f"=SUM('Monthly Model'!AC{start+5}:AC{end+5})",
                       f"=SUM('Monthly Model'!AC6:AC{end+5})",
                       f"='Monthly Model'!AP{end+5}",
                       f"=MAX('Monthly Model'!AK6:AK{end+5})"])
            for col in range(4, 8):
                ws.cell(ws.max_row, col).number_format = MONEY

    def staff_months(self):
        ws = self.sheet("Staff Months", "Staffing monthly calculation detail",
                        "Formula audit trail: date-bounded capacity; positive incremental FTE triggers recruiting/equipment. Only employees receive wage-based burden. Fractional headcount slots are ceiling(FTE), not confirmed people.",
                        ["Scenario", "Role ID", "Month", "Date", "FTE", "Salary cash", "Employer burden cash", "Loaded cash", "Unpaid founder economic", "Recruiting", "Equipment", "Headcount slots"])
        for scenario in SCENARIOS:
            for staff in self.data["staff"]:
                sr = self.staff_rows[(scenario, staff["id"])]
                s = lambda col: f"'Staffing'!${col}${sr}"
                scope = staff["scope"]
                start = s("F") if scope == "Core" else self.a(scope.lower() + "_start", scenario)
                for month in range(1, 61):
                    r = ws.max_row + 1
                    enabled = self.scope_formula(scope)
                    fte = f'=IF(AND(C{r}>={start},C{r}<={s("G")},{enabled}=1),IF(C{r}<=6,{s("H")},IF(C{r}<=15,{s("I")},IF(C{r}<=30,{s("J")},{s("K")}))),0)'
                    salary = self.a("founder_salary", scenario) if staff["type"] == "Founder" else s("L")
                    economic = f'E{r}*{salary}/12*(1+{self.a("salary_growth", scenario)})^INT((C{r}-1)/12)'
                    paid = "'Start Here'!$B$10" if staff["type"] == "Founder" else "1"
                    previous = f"E{r-1}" if month > 1 else "0"
                    ws.append([scenario, staff["id"], month,
                               f"=EDATE('Start Here'!$B$6,C{r}-1)", fte,
                               f"={economic}*{paid}", f"=F{r}*{s('N')}", f"=F{r}+G{r}",
                               f"={economic}*(1-{paid})" if staff["type"] == "Founder" else "=0",
                               f"=MAX(0,E{r}-{previous})*{s('O')}",
                               f"=MAX(0,E{r}-{previous})*{s('P')}", f"=ROUNDUP(E{r},0)"])
                    ws.cell(r, 4).number_format = "mmm yyyy"
                    for col in range(6, 12):
                        ws.cell(r, col).number_format = MONEY
        ws.column_dimensions.group("F", "K", hidden=False)

    def monthly(self):
        headers = ["Month", "Date", "New customers", "Active customers", "Users", "API requests M", "Storage GB", "Logs GB",
                   "Total FTE", "Headcount slots", "Staff cash", "Recruiting", "Equipment",
                   *CATEGORIES, "Vendor total", "Contingency", "Payment fees", "Total spend", "Monthly revenue",
                   "Exit ARR", "Net burn", "Financing", "Opening cash", "Closing cash", "Cumulative spend",
                   "Funding gap", "Cash if zero revenue", "Unresolved vendor cash", "Optional vendor cash",
                   "Unpaid founder economic", "Cumulative net need", "Cost-only burn", "Negative flag"]
        for scenario in SCENARIOS:
            ws = self.sheet("Monthly " + scenario, f"{scenario} | 60-month financial model",
                            "Expected-value cohorts, monthly revenue recognized = active organizations × ARPA; ARR = exit-month revenue ×12, never annual revenue. Spend includes contingency; cash excludes taxes, financing fees and receivables timing.",
                            headers)
            sm = self.wb["Staff Months"].max_row
            vm = 5 + len(SCENARIOS) * 60 * len(self.data["costs"])
            for month in range(1, 61):
                r = month + 5
                prev = r - 1
                sum_staff = lambda col: f'SUMIFS(\'Staff Months\'!${col}$6:${col}${sm},\'Staff Months\'!$A$6:$A${sm},$AS$1,\'Staff Months\'!$C$6:$C${sm},A{r})'
                sum_vendor = lambda col, category=None: f'SUMIFS(\'Vendor Months\'!${col}$6:${col}${vm},\'Vendor Months\'!$A$6:$A${vm},$AS$1,\'Vendor Months\'!$C$6:$C${vm},A{r}' + (f',\'Vendor Months\'!$J$6:$J${vm},{category}' if category else "") + ")"
                ws.append([
                    month, f"=EDATE('Start Here'!$B$6,A{r}-1)",
                    f'=IF(A{r}>={self.a("launch",scenario)},{self.a("new_customers",scenario)}*(1+{self.a("customer_growth",scenario)})^INT((A{r}-{self.a("launch",scenario)})/12),0)',
                    f'={(f"D{prev}" if month>1 else "0")}*(1-{self.a("churn",scenario)})+C{r}',
                    f'=D{r}*{self.a("users",scenario)}', f'=E{r}*{self.a("requests",scenario)}*30/1000000',
                    f'=D{r}*{self.a("storage",scenario)}', f'=D{r}*{self.a("logs",scenario)}',
                    "=" + sum_staff("E"), "=" + sum_staff("L"), "=" + sum_staff("H"),
                    "=" + sum_staff("J"), "=" + sum_staff("K"),
                    *["=" + sum_vendor("G", '"' + cat + '"') for cat in CATEGORIES],
                    f"=SUM(N{r}:Y{r})", f'=(SUM(K{r}:M{r})+Z{r})*{self.a("contingency",scenario)}',
                    f'=AD{r}*{self.a("payment_rate",scenario)}',
                    f"=SUM(K{r}:M{r})+Z{r}+AA{r}+AB{r}",
                    f'=D{r}*{self.a("arpa",scenario)}', f"=AD{r}*12", f"=AC{r}-AD{r}",
                    f'=IF(A{r}={self.a("financing_month",scenario)},{self.a("financing_amount",scenario)},0)',
                    f'={self.a("initial_cash",scenario)}' if month == 1 else f"=AI{prev}",
                    f"=AH{r}+AG{r}-AF{r}",
                    f"=SUM($AC$6:AC{r})", f"=MAX(0,-AI{r})",
                    f'={self.a("initial_cash",scenario)}+SUM($AG$6:AG{r})-AJ{r}',
                    "=" + sum_vendor("H"), f"=X{r}+Y{r}", "=" + sum_staff("I"),
                    f"=SUM($AF$6:AF{r})", f"=AC{r}", f"=IF(AI{r}<0,1,0)"
                ])
                for col in range(3, 45):
                    ws.cell(r, col).number_format = NUMBER if col < 11 else MONEY
                ws.cell(r, 2).number_format = "mmm yyyy"
                ws.cell(r, 44).number_format = "0"
            ws["AS1"] = scenario
            ws.column_dimensions["AS"].hidden = True
            ws.column_dimensions.group("E", "H", hidden=True)
            ws.column_dimensions.group("L", "Y", hidden=True)
            ws.column_dimensions.group("AJ", "AR", hidden=True)
        ws = self.sheet("Monthly Model", "Selected scenario | monthly cash and workload",
                        "Select scenario on Start Here. Grouped columns expand to show workload/category detail. Financing is hypothetical. Cost-only cash keeps forecast usage/costs but removes revenue receipts (conservative sensitivity).",
                        headers)
        for r in range(6, 66):
            for col in range(1, 45):
                c = letter(col)
                ws.cell(r, col, f'=IF(\'Start Here\'!$B$5="Lean",\'Monthly Lean\'!{c}{r},IF(\'Start Here\'!$B$5="Growth",\'Monthly Growth\'!{c}{r},\'Monthly Base\'!{c}{r}))')
                ws.cell(r, col).number_format = "mmm yyyy" if col == 2 else NUMBER if col <= 10 else MONEY
        ws.column_dimensions.group("E", "H", hidden=True)
        ws.column_dimensions.group("L", "Y", hidden=True)
        ws.column_dimensions.group("AJ", "AR", hidden=True)

    def vendor_months(self):
        ws = self.sheet("Vendor Months", "Vendor monthly quantity × rate audit trail",
                        "Unknown prices are explicit allowances, never missing-as-free. Stack and scope gates applied before quantity. Inflation steps on project anniversaries. Infrastructure allowances include dev/stage/prod; not audited COGS.",
                        ["Scenario", "Cost ID", "Month", "Date", "Enabled quantity", "Escalated unit rate", "Cash", "Unresolved cash", "Optional cash", "Category"])
        for scenario in SCENARIOS:
            m = "'Monthly " + scenario + "'!"
            for cost in self.data["costs"]:
                cr = self.cost_rows[cost[0]]
                n = lambda col: f"'Nonlabor'!${col}${cr}"
                for month in range(1, 61):
                    r, mr = ws.max_row + 1, month + 5
                    ai_start, ct_start = self.a("ai_start", scenario), self.a("ctms_start", scenario)
                    start = f'IF({n("E")}="Core",{n("F")},MAX({n("F")},IF({n("E")}="AI",{ai_start},{ct_start})))'
                    scope = f'IF({n("E")}="Core",1,IF({n("E")}="AI",\'Start Here\'!$B$7,\'Start Here\'!$B$8))'
                    quantity = "0"
                    drivers = {"Fixed": "1", "Seats": f"ROUNDUP({m}I{mr},0)", "Customers": f"{m}D{mr}",
                               "NewCustomers": f"{m}C{mr}", "StorageGB": f"{m}G{mr}", "LogGB": f"{m}H{mr}",
                               "RequestsM": f"{m}F{mr}",
                               "InputM": f'{m}D{mr}*{self.a("tokens_in",scenario)}/1000000',
                               "OutputM": f'{m}D{mr}*{self.a("tokens_out",scenario)}/1000000'}
                    for driver, expr in reversed(list(drivers.items())):
                        quantity = f'IF({n("I")}="{driver}",{expr},{quantity})'
                    occurs = f'OR({n("H")}="Monthly",AND({n("H")}="Once",C{r}={start}),AND({n("H")}="Annual",MOD(C{r}-{start},12)=0))'
                    enabled = f'AND(OR({n("D")}="Both",{n("D")}=\'Start Here\'!$B$9),{scope}=1,C{r}>={start},C{r}<={n("G")},{occurs})'
                    ws.append([scenario, cost[0], month, f"=EDATE('Start Here'!$B$6,C{r}-1)",
                               f"=IF({enabled},{quantity},0)",
                               f'={n("J")}*(1+{self.a("inflation",scenario)})^INT((C{r}-1)/12)',
                               f"=E{r}*F{r}",
                               f'=IF(OR({n("L")}="Quote required",{n("L")}="Unverified"),G{r},0)',
                               f'=IF({n("E")}="Core",0,G{r})', f'={n("B")}'])
                    ws.cell(r, 4).number_format = "mmm yyyy"
                    for col in (6, 7, 8, 9):
                        ws.cell(r, col).number_format = MONEY

    def summaries(self):
        ws = self.sheet("Annual Summary", "Annual reconciliation | selected scenario",
                        "Year 1–5 = consecutive 12-month project years. Calendar summaries include partial first/last years; both sum to the same 60-month spend. Cash is closing balance, not sum of balances.",
                        ["Basis", "Year", "Start date", "End date", "Spend", "Revenue", "Net burn", "Financing", "Closing cash", "Peak FTE", "Staff cash", "Infra", "Tools", "Other vendor", "One-time recruitment/equipment", "Contingency", "Payment fees"])
        for year in range(1, 6):
            r = year + 5
            start, end = (year - 1) * 12 + 6, year * 12 + 5
            ws.append(["Project", year, f"='Monthly Model'!B{start}", f"='Monthly Model'!B{end}",
                       *[f"=SUM('Monthly Model'!{c}{start}:{c}{end})" for c in ("AC", "AD", "AF", "AG")],
                       f"='Monthly Model'!AI{end}", f"=MAX('Monthly Model'!I{start}:I{end})",
                       f"=SUM('Monthly Model'!K{start}:K{end})", f"=SUM('Monthly Model'!N{start}:N{end})",
                       f"=SUM('Monthly Model'!O{start}:O{end})", f"=SUM('Monthly Model'!P{start}:Y{end})",
                       f"=SUM('Monthly Model'!L{start}:M{end})", f"=SUM('Monthly Model'!AA{start}:AA{end})",
                       f"=SUM('Monthly Model'!AB{start}:AB{end})"])
        # Sixty periods span at most six calendar years.
        for offset in range(6):
            r = ws.max_row + 1
            ws.append(["Calendar", "=YEAR('Start Here'!$B$6)" if offset == 0 else f"=B{r-1}+1",
                       f"=DATE(B{r},1,1)", f"=DATE(B{r},12,31)"])
            for col, source in [(5, "AC"), (6, "AD"), (7, "AF"), (8, "AG"), (11, "K"), (12, "N"), (13, "O"), (16, "AA"), (17, "AB")]:
                ws.cell(r, col, f'=SUMIFS(\'Monthly Model\'!{source}6:{source}65,\'Monthly Model\'!$B$6:$B$65,">="&C{r},\'Monthly Model\'!$B$6:$B$65,"<="&D{r})')
            ws.cell(r, 9, f'=IF(SUMIFS(\'Monthly Model\'!A6:A65,\'Monthly Model\'!B6:B65,">="&C{r},\'Monthly Model\'!B6:B65,"<="&D{r})=0,0,INDEX(\'Monthly Model\'!AI6:AI65,COUNTIF(\'Monthly Model\'!B6:B65,"<="&D{r})))')
        for r in range(6, ws.max_row + 1):
            ws.cell(r, 3).number_format = ws.cell(r, 4).number_format = "yyyy-mm-dd"
            for col in range(5, 18):
                ws.cell(r, col).number_format = MONEY if col != 10 else NUMBER
        ws["A25"] = "Reconciliation difference (should be 0)"
        ws["E25"] = "=SUM(E6:E10)-SUM('Monthly Model'!AC6:AC65)"
        ws["E26"] = "=SUM(E11:E16)-SUM(E6:E10)"

        ws = self.sheet("Scenario Comparison", "Scenario comparison | identical options, different operating assumptions",
                        "All three scenarios are live formulas, same global scope/stack switches. Comparison is spend, not valuation. Revenue is unproven; cost-only funding gap is separately visible.",
                        ["Scenario", "5-year spend", "Year 1 spend", "Peak FTE", "Cash low", "Funding gap", "Cost-only gap", "Unresolved vendor cash", "Unpaid founder", "Year 5 exit ARR"])
        for scenario in SCENARIOS:
            m = "'Monthly " + scenario + "'!"
            ws.append([scenario, f"=SUM({m}AC6:AC65)", f"=SUM({m}AC6:AC17)", f"=MAX({m}I6:I65)",
                       f"=MIN({m}AI6:AI65)", f"=MAX({m}AK6:AK65)", f"=MAX(0,-MIN({m}AL6:AL65))",
                       f"=SUM({m}AM6:AM65)", f"=SUM({m}AO6:AO65)", f"={m}AE65"])
            for col in range(2, 11):
                ws.cell(ws.max_row, col).number_format = NUMBER if col == 4 else MONEY

        ws = self.sheet("Charts", "Formula-linked planning charts",
                        "Charts link to financial cells; no screenshots or static chart values. Expand monthly category columns to audit.")
        self.chart(ws, "Monthly Model", "Monthly cash spend and revenue", [29, 30], 6, 65, 2, "A5")
        self.chart(ws, "Monthly Model", "Closing cash / zero-revenue sensitivity", [35, 38], 6, 65, 2, "K5")
        self.chart(ws, "Monthly Model", "FTE ramp incl. founder", [9], 6, 65, 2, "A21")
        self.chart(ws, "Monthly Model", "Cumulative spend and net cash needed", [36, 42], 6, 65, 2, "K21")
        self.chart(ws, "Annual Summary", "Annual cost categories", [11, 12, 13, 14, 15, 16, 17], 6, 10, 2, "A37", stacked=True)
        self.chart(ws, "Scenario Comparison", "Five-year scenario spend", [2], 6, 8, 1, "K37", bar=True)
        self.chart(self.wb["Start Here"], "Annual Summary", "Project-year cash spend", [5], 6, 10, 2, "D12", bar=True)

    def chart(self, dest, sheet, title, columns, start, end, category_col, anchor, stacked=False, bar=False):
        chart = BarChart() if bar or stacked else LineChart()
        chart.title = title
        chart.style = 13
        chart.width, chart.height = 22, 8
        if stacked:
            chart.grouping = "stacked"
            chart.overlap = 100
        source = self.wb[sheet]
        for col in columns:
            chart.add_data(Reference(source, min_col=col, min_row=start - 1, max_row=end), titles_from_data=True)
        chart.set_categories(Reference(source, min_col=category_col, min_row=start, max_row=end))
        dest.add_chart(chart, anchor)

    def finish(self):
        for ws in self.wb:
            if ws.title != "Start Here":
                ws.auto_filter.ref = f"A5:{letter(ws.max_column)}{ws.max_row}"
                for col in range(1, ws.max_column + 1):
                    dim = ws.column_dimensions[letter(col)]
                    if dim.width == 13:
                        dim.width = 20
            for row in ws.iter_rows(min_row=6):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.data_type == "f":
                        cell.fill = PatternFill("solid", fgColor=GREEN)
                    elif cell.fill.fgColor.rgb == "00000000":
                        cell.font = Font(size=10, color="333333")
            if ws.title.startswith("Monthly"):
                for c in ("AI", "AL"):
                    ws.conditional_formatting.add(f"{c}6:{c}65", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE")))
            ws.print_title_rows = "1:5"
        start = self.wb["Start Here"]
        for r, name in enumerate(self.wb.sheetnames[1:], 31):
            start.cell(r, 1, name).hyperlink = f"#'{name}'!A1"
            start.cell(r, 1).style = "Hyperlink"
        start["D5"] = "Workbook navigation ↓ (also see tabs)"
        start["D6"] = "Source refresh and approvals required before commitments."
        start["D7"] = "No macros, external workbook links, or network calls in generator."
        start["D31"] = "Five-year cash by category"
        dashboard_categories = [("Staff cash", "K"), ("Recruiting", "L"), ("Equipment", "M")]
        dashboard_categories.extend(zip(CATEGORIES, (letter(c) for c in range(14, 26))))
        dashboard_categories.extend([("Contingency", "AA"), ("Payment fees", "AB")])
        for r, (label, col) in enumerate(dashboard_categories, 32):
            start.cell(r, 4, label)
            start.cell(r, 5, f"=SUM('Monthly Model'!{col}6:{col}65)").number_format = MONEY
            start.cell(r, 5).fill = PatternFill("solid", fgColor=GREEN)
        start["D51"] = "Optional staffing is in Staff cash, not in optional vendor-only categories."
        start["D52"] = "WBS assumed person-days use these staffing pools; not additional project labor invoices."

    def build(self):
        self.setup()
        self.catalogs()
        self.traceability()
        self.staff_months()
        self.monthly()
        self.vendor_months()
        self.summaries()
        self.finish()
        return self.wb


def snapshot(data):
    base = {s: evaluate(data, s) for s in SCENARIOS}
    sensitivity = {}
    for name, ai, ctms in (("AI only", 1, 0), ("CTMS only", 0, 1), ("AI and CTMS", 1, 1)):
        changed = copy.deepcopy(data)
        changed["controls"].update(ai=ai, ctms=ctms)
        sensitivity[name] = {}
        for scenario in SCENARIOS:
            result = evaluate(changed, scenario)
            sensitivity[name][scenario] = {k: result[k] for k in ("spend", "peak_fte", "gap")}
    return {
        "basis": "Python computed companion, not Excel cached values; regenerate after input JSON changes. Workbook edits require Excel recalculation and are not imported back.",
        "start": data["start"], "controls": data["controls"], "scenarios": base,
        "option_sensitivity": sensitivity
    }


def headline_markdown(computed):
    lines = [
        "| Scenario | Five-year cash spend | Year 1 spend | Peak FTE incl. founder | Cash low | Additional funding gap |",
        "|---|---:|---:|---:|---:|---:|"
    ]
    for scenario, result in computed["scenarios"].items():
        lines.append(f"| {scenario} | ${result['spend']:,.0f} | ${result['year1']:,.0f} | {result['peak_fte']:.2f} | ${result['cash_low']:,.0f} | ${result['gap']:,.0f} |")
    lines.extend(["", "Optional scope sensitivity (total five-year spend; incremental staffing and vendor cash included):", "",
                  "| Enabled options | Lean | Base | Growth |", "|---|---:|---:|---:|"])
    for name, scenarios in computed["option_sensitivity"].items():
        lines.append("| " + name + " | " + " | ".join(f"${scenarios[s]['spend']:,.0f}" for s in SCENARIOS) + " |")
    return "\n".join(lines)


def recalculate(path):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("Recalculation requires installed LibreOffice; use recalc-on-open otherwise")
    with tempfile.TemporaryDirectory(prefix="saas-recalc-") as temp:
        temp_path = Path(temp)
        converted = temp_path / "converted"
        converted.mkdir()
        subprocess.run([
            executable, f"-env:UserInstallation={temp_path.as_uri()}/profile",
            "--headless", "--convert-to", "xlsx", "--outdir", str(converted), str(path)
        ], check=True, timeout=240, capture_output=True)
        cached = converted / path.name
        if not cached.is_file():
            raise RuntimeError("LibreOffice did not produce a recalculated workbook")
        shutil.copyfile(cached, path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT)
    parser.add_argument("--recalculate", action="store_true",
                        help="Use installed LibreOffice headless to populate Excel formula caches")
    args = parser.parse_args()
    data = load_inputs()
    scope = json.loads((HERE / "scope.json").read_text())
    sources = json.loads((HERE / "sources.json").read_text())
    args.output_dir.mkdir(parents=True, exist_ok=True)
    wb = Planner(data, scope, sources).build()
    path = args.output_dir / "Clinical-Roundup-SaaS-Build-and-Scale.xlsx"
    wb.save(path)
    if args.recalculate:
        recalculate(path)
    computed = snapshot(data)
    (args.output_dir / "snapshot.json").write_text(json.dumps(computed, indent=2) + "\n")
    readme = args.output_dir / "README.md"
    if readme.exists():
        updated = re.sub(r"(?s)(<!-- computed-results:start -->\n).*?(\n<!-- computed-results:end -->)",
                         lambda match: match[1] + headline_markdown(computed) + match[2],
                         readme.read_text())
        readme.write_text(updated)
    for scenario, result in computed["scenarios"].items():
        print(f"{scenario}: 5yr=${result['spend']:,.0f}; Y1=${result['year1']:,.0f}; peak FTE={result['peak_fte']:.2f}; low cash=${result['cash_low']:,.0f}; gap=${result['gap']:,.0f}")
    print(path)


if __name__ == "__main__":
    main()
